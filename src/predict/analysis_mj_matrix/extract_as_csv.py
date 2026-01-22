import pandas as pd
import numpy as np


pdb_id = '6axz'

sequence = ''.join(np.load(f'../../../data/FormatData/{pdb_id}.npz', allow_pickle=True)['sequence'])

vqe_matrix = np.load('./vqe/interactions.npz', allow_pickle=True)[pdb_id]
for i in range(vqe_matrix.shape[0]):
    for j in range(i+1, vqe_matrix.shape[0]):
        vqe_matrix[j, i] = vqe_matrix[i, j]
vqe_matrix = -vqe_matrix

qnn_matrix = np.load('./qnn/interactions.npz', allow_pickle=True)[pdb_id].item()
qnn_matrix = np.stack([qnn_matrix[key] for key in range(5)], axis=0).mean(axis=0)
qnn_matrix[np.eye(qnn_matrix.shape[0], dtype=bool)] = 0

# for vqe
for i in range(vqe_matrix.shape[0]):
    for j in range(vqe_matrix.shape[1]):
        if np.isclose(vqe_matrix[i, j], 0):
            vqe_matrix[i, j] = np.nan

# for qnn
qnn_matrix[np.eye(qnn_matrix.shape[0], dtype=bool)] = np.nan
qnn_matrix[0, :] = np.nan

def scale_matrix(matrix):
    mask = np.ones_like(matrix, dtype=bool)
    for i in range(matrix.shape[0]):
        for j in range(matrix.shape[1]):
            if np.isnan(matrix[i, j]):
                mask[i, j] = False
    
    value_matrix = matrix[mask]
    value_matrix = (value_matrix - value_matrix.min()) / (value_matrix.max() - value_matrix.min())
    matrix[mask] = value_matrix
    return matrix

qnn_matrix = scale_matrix(qnn_matrix)
vqe_matrix = scale_matrix(vqe_matrix)

from openpyxl import Workbook
wb = Workbook()

def build(sheet_name, matrix):
    sheet = wb.create_sheet(sheet_name)
    for i in range(len(sequence)+1):
        for j in range(len(sequence)+1):
            if i==0 and j==0:
                continue
            elif i==0:
                sheet.cell(row=i+1, column=j+1, value=sequence[j-1])
            elif j==0:
                sheet.cell(row=i+1, column=j+1, value=sequence[i-1])
            else:
                sheet.cell(row=i+1, column=j+1, value=matrix[i-1, j-1])

# print(vqe_matrix)
# print(qnn_matrix)
build('VQE', vqe_matrix)
build('QNN', qnn_matrix)

wb.save(f'./{pdb_id}.xlsx')

# # plot
# from matplotlib import pyplot as plt
# import seaborn as sns

# fig, ax = plt.subplots(1, 2, figsize=(20, 10))

# sns.heatmap(vqe_matrix, ax=ax[0], cmap='coolwarm', xticklabels=sequence, yticklabels=sequence)
# sns.heatmap(qnn_matrix, ax=ax[1], cmap='coolwarm', xticklabels=sequence, yticklabels=sequence)

# plt.show()